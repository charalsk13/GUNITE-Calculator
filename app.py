from io import BytesIO
from pathlib import Path
import importlib.util

import pandas as pd
import streamlit as st

import catalog
from catalog import available_diameters, available_spacings
from gunite_calculator import GuniteInput, calculate, order_rows, result_table
import database
import project_io
from pdf_report import create_project_pdf

try:
    import drawing
except ImportError:
    drawing_path = Path(__file__).with_name("drawing.py")
    spec = importlib.util.spec_from_file_location("drawing_local", drawing_path)
    if spec is None or spec.loader is None:
        raise
    drawing_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(drawing_module)
    drawing = drawing_module


st.set_page_config(page_title="GUNITE | Υπολογισμός & Παραγγελία", layout="wide")

st.markdown(
    """
    <style>
    :root { --gunite-accent:#2563eb; --gunite-ink:#172033; --gunite-muted:#667085; --gunite-border:#e5e7eb; --gunite-soft:#f6f8fb; }
    .block-container {max-width: 1480px; padding-top: 1.25rem; padding-bottom: 4rem;}
    .hero {padding: 1.55rem 1.65rem; border: 1px solid var(--gunite-border); border-radius: 22px;
           background: linear-gradient(135deg,#ffffff 0%,#f4f7fb 100%); margin-bottom: 1.2rem;
           box-shadow: 0 10px 30px rgba(16,24,40,.06);}
    .hero-brand {display:flex; align-items:center; gap:.85rem;}
    .hero-mark {width:42px;height:42px;border-radius:12px;background:#172033;color:#fff;display:flex;align-items:center;justify-content:center;font-weight:800;letter-spacing:.5px;}
    .hero h1 {margin:0; font-size:2.05rem; letter-spacing:-.035em; color:var(--gunite-ink);}
    .hero .subtitle {margin:.28rem 0 0 0; color:var(--gunite-muted); font-size:.96rem;}
    .hero .context {margin-top:1rem;display:flex;gap:.55rem;flex-wrap:wrap;}
    .chip {display:inline-flex;align-items:center;padding:.35rem .65rem;border:1px solid #dfe4ec;border-radius:999px;background:#fff;color:#475467;font-size:.78rem;font-weight:650;}
    .section-kicker {font-size:.74rem;text-transform:uppercase;letter-spacing:.11em;font-weight:800;color:#667085;margin-bottom:.25rem;}
    .section-title {font-size:1.55rem;font-weight:800;letter-spacing:-.025em;color:var(--gunite-ink);margin:.05rem 0 .25rem;}
    .section-help {color:var(--gunite-muted);font-size:.9rem;margin-bottom:1rem;}
    .card {padding:1rem 1.05rem; border:1px solid var(--gunite-border); border-radius:15px; background:#fff; min-height:82px; box-shadow:0 4px 14px rgba(16,24,40,.035);}
    .card .label {font-size:.74rem; color:var(--gunite-muted); margin-bottom:.28rem; font-weight:650;}
    .card .value {font-size:1.18rem; font-weight:800; color:var(--gunite-ink);}
    .source-card {padding:.9rem 1rem;border:1px solid #dbe4f0;border-radius:14px;background:#f8fafc;}
    .source-card strong {color:#344054;}
    .source-card span {color:#667085;font-size:.84rem;}
    .workflow {display:flex;gap:.55rem;align-items:center;margin:.2rem 0 1.35rem;}
    .workflow-item {flex:1;min-width:120px;padding:.75rem .8rem;border:1px solid var(--gunite-border);border-radius:13px;background:#fff;}
    .workflow-item .num {font-size:.7rem;font-weight:800;color:#98a2b3;letter-spacing:.08em;}
    .workflow-item .name {font-size:.88rem;font-weight:750;color:#475467;margin-top:.15rem;}
    .workflow-item.active {border-color:#93b4f5;background:#eff6ff;box-shadow:inset 0 0 0 1px #dbeafe;}
    .workflow-item.active .num,.workflow-item.active .name {color:#1d4ed8;}
    .workflow-arrow {color:#98a2b3;font-size:1rem;}
    .stRadio > div {gap:.45rem;}
    .stRadio label {border:1px solid var(--gunite-border);border-radius:12px;padding:.42rem .72rem;background:#fff;}
    div[data-testid="stSidebar"] {border-right:1px solid #e7eaf0;}
    div[data-testid="stSidebar"] .block-container {padding-top:1.1rem;}
    .sidebar-title {font-weight:800;font-size:1.08rem;color:var(--gunite-ink);}
    .sidebar-note {font-size:.78rem;color:#667085;line-height:1.45;}
    .step-note {padding:.75rem .9rem;border-left:3px solid #2563eb;background:#f5f8ff;border-radius:0 10px 10px 0;color:#475467;font-size:.86rem;}
    .ok-card {padding:1rem 1.1rem;border:1px solid #b7dfc5;border-radius:14px;background:#f2fbf5;}
    .warn-card {padding:1rem 1.1rem;border:1px solid #ead59b;border-radius:14px;background:#fffaf0;}
    .small-note {font-size:.82rem;color:var(--gunite-muted);}
    div.stButton > button, div.stDownloadButton > button {border-radius:11px;font-weight:700;min-height:2.55rem;}
    div.stButton > button[kind="primary"] {box-shadow:0 6px 14px rgba(37,99,235,.16);}
    </style>
    """,
    unsafe_allow_html=True,
)


def dataframe_to_excel_bytes(sheets: dict, title: str = "GUNITE") -> bytes:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = str(sheet_name)[:31]
            df = pd.DataFrame() if df is None else df
            df.to_excel(writer, sheet_name=safe_name, index=False, startrow=2)
            ws = writer.sheets[safe_name]
            ws["A1"] = title
            ws["A1"].font = Font(bold=True, size=14)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))
            if ws.max_row >= 3:
                for cell in ws[3]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9EAF7")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                thin = Side(style="thin", color="B7B7B7")
                for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
                    for cell in row:
                        cell.border = Border(bottom=thin)
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.freeze_panes = "A4"
                ws.auto_filter.ref = ws.dimensions
            # Do not use column_cells[0].column_letter here: after merging A1:...,
            # the first cell of columns B..N is an openpyxl MergedCell and has no
            # column_letter attribute. Use the numeric column index instead.
            for col_idx, column_cells in enumerate(ws.iter_cols(), start=1):
                max_length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 3, 45)
    return output.getvalue()


def input_dataframe(inp):
    is_4p = (inp.x1b > 0 and inp.x2b > 0 and inp.y1b > 0 and inp.y2b > 0
              and abs(inp.x1b - inp.x2b) < 1e-9 and abs(inp.y1b - inp.y2b) < 1e-9)
    total_x = max(inp.x1b, inp.x2b) if is_4p else inp.x1b + inp.x2b
    total_y = max(inp.y1b, inp.y2b) if is_4p else inp.y1b + inp.y2b
    rows = [("Στοιχείο", inp.name),
            ("X1b (cm) — είσοδος DOCX", inp.x1b), ("X2b (cm) — είσοδος DOCX", inp.x2b),
            ("Y1b (cm) — είσοδος DOCX", inp.y1b), ("Y2b (cm) — είσοδος DOCX", inp.y2b),
            ("Υφιστάμενη X (cm) — φυσική διάσταση", total_x), ("Υφιστάμενη Y (cm) — φυσική διάσταση", total_y),
            ("Φd — βασική (mm)", inp.fd_mm), ("Φd — γωνίες (mm)", inp.corner_fd_mm or inp.fd_mm),
            ("Φd — πλευρές X (mm)", inp.x_side_fd_mm or inp.fd_mm), ("Φd — πλευρές Y (mm)", inp.y_side_fd_mm or inp.fd_mm),
            ("Φs (mm)", inp.fs_mm), ("s (mm)", inp.spacing_mm), ("a1 (cm)", inp.a1_cm),
            ("a2 (cm)", inp.a2_cm), ("tgun (cm)", inp.tgun_cm), ("dh (cm)", inp.dh_cm),
            ("Ύψος ορόφου (m)", inp.floor_height_m), ("Αναμονή (m)", inp.waiting_m),
            ("nx", inp.nx), ("ny", inp.ny), ("Τρόπος ανάπτυξης 4Π", getattr(inp, "development_option", "1a")), ("Συντελεστής υπερκάλυψης ×Φd", inp.overlap_factor)]
    return pd.DataFrame(rows, columns=["Παράμετρος", "Τιμή"])


def calculation_dataframe(inp, res):
    rows = [("Περίπτωση", res.case), ("d (cm)", res.d_cm),
            ("Υφιστάμενη διατομή (cm)", f"{res.existing_width_m*100:.1f} × {res.existing_height_m*100:.1f}"),
            ("Τελική gunite (cm)", f"{res.gunite_width_m*100:.1f} × {res.gunite_height_m*100:.1f}"),
            ("X1g (cm)", res.x1g), ("X2g (cm)", res.x2g), ("Y1g (cm)", res.y1g), ("Y2g (cm)", res.y2g),
            ("X1s (cm)", res.x1s), ("X2s (cm)", res.x2s), ("Y1s (cm)", res.y1s), ("Y2s (cm)", res.y2s),
            ("Περίμετρος gunite (m)", res.gunite_perimeter_m), ("Επιφάνεια gunite (m²)", res.gunite_surface_m2),
            ("Όγκος gunite (m³)", res.gunite_volume_m3), ("Διαμήκεις ράβδοι", f"{res.longitudinal_bars_total} Φ{inp.fd_mm:.0f}"),
            ("Μήκος διαμήκους (m)", res.rounded_main_bar_length_m), ("Υπερκάλυψη (m)", res.overlap_length_m)]
    return pd.DataFrame(rows, columns=["Υπολογισμός", "Αποτέλεσμα"])


def jacket_dataframe(inp, res):
    rows = []
    specs = [
        (res.p_development.label, res.catalog_name, res.catalog_weight_kg),
    ]
    if res.p_development_bottom is not None:
        specs.append((res.p_development_bottom.label, res.catalog_name_bottom, res.catalog_weight_kg_bottom))
    for position, name, weight in specs:
        rows.append({
            "Θέση": position,
            "Γεωμετρικό (cm)": round((res.p_development if position == res.p_development.label else res.p_development_bottom).geometric_required_m * 100, 1),
            "Με συγκόλληση (cm)": round((res.p_development if position == res.p_development.label else res.p_development_bottom).required_m * 100, 1),
            "Προϊόν καταλόγου": name or "Δεν βρέθηκε",
            "Τεμάχια": 1,
            "Υπόλοιπο/άκρο (cm)": round((res.p_development if position == res.p_development.label else res.p_development_bottom).surplus_each_end_cm, 1),
            "Βάρος/τεμ. (kg)": weight,
            "Σύνολο kg": weight,
        })
    return pd.DataFrame(rows)



def element_schedule_row(item, local_index: int):
    """Return one row matching the user's reference table for a project element."""
    inp, res = item["input"], item["result"]
    case_num = "4" if res.case.startswith("4") else "3" if "3πλευρος" in res.case else "2"
    # The reference table shows actual section dimensions separately from study dimensions.
    real_dims = f"{item.get('length_cm', inp.x1b + inp.x2b):.3f}/{item.get('width_cm', inp.y1b + inp.y2b):.3f}"
    # Counts by actual selected longitudinal diameter. This supports mixed bars.
    diameter_counts = {}
    for info in res.longitudinal_groups.values():
        count = int(info.get("count", 0))
        dia = int(round(float(info.get("diameter_mm", inp.fd_mm))))
        if count > 0:
            diameter_counts[dia] = diameter_counts.get(dia, 0) + count

    # For the reference-style table, X1/Y1/X2/Y2 are reported in metres as
    # generated geometric values. They are not user inputs.
    x1 = res.x1s / 100.0
    y1 = res.y1s / 100.0
    x2 = res.x2s / 100.0
    y2 = res.y2s / 100.0
    pieces = 2 if res.p_development_bottom is not None else 1
    # Keep the catalogue pair visible in the row without inventing a value.
    jacket_size = f"{inp.fs_mm:.0f}/{inp.spacing_mm}"
    return {
        "α/α": local_index,
        "Στοιχείο": inp.name,
        "Διαστάσεις Μελέτης": item.get("study_dimensions", "") or "—",
        "Διαστάσεις πραγματικές": real_dims,
        "ΠΛΕΥΡΕΣ ΜΑΝΔΥΑ": case_num,
        "ΠΑΡΑΤΗΡΗΣΕΙΣ": item.get("remarks", "") or "",
        "ΠΑΡΑΤΗΡΗΣΕΙΣ ΜΕΛΕΤΗΤΗ": item.get("designer_remarks", "") or "",
        "ΟΠΛΙΣΜΟΣ ΔΙΑΜΕΤΡΟΣ Φ": f"Φ{inp.fd_mm:.0f}",
        "ΟΠΛΙΣΜΟΙ ΔΙΑΜΕΤΡΟΥ Φ20": diameter_counts.get(20, 0),
        "ΟΠΛΙΣΜΟΙ ΔΙΑΜΕΤΡΟΥ Φ25": diameter_counts.get(25, 0),
        "Κατακόρυφα Φ20": "",
        "Κατακόρυφα Φ25": "",
        "ΥΨΟΣ ΟΡΟΦΟΥ Y": f"{inp.floor_height_m:.3f}",
        "ΥΨΟΣ ΒΕΡΓΑΣ": f"{res.rounded_main_bar_length_m:.2f}",
        "ΔΙΑΤΟΜΗ ΜΑΝΔΥΑ/ΑΠΟΣΤΑΣΗ": jacket_size,
        "X1": f"{x1:.2f}",
        "Y1": f"{y1:.2f}",
        "X2": f"{x2:.2f}",
        "Y2": f"{y2:.2f}",
        "Ύψος για Μανδύα": f"{inp.floor_height_m:.2f}",
        "ΤΕΜΑΧΙΑ ΜΑΝΔΥΑ": pieces,
    }

def build_element_excel(inp, res):
    return dataframe_to_excel_bytes({"Είσοδοι":input_dataframe(inp), "Υπολογισμοί":calculation_dataframe(inp,res),
                                     "Μανδύες":jacket_dataframe(inp,res), "Παραγγελία":pd.DataFrame(order_rows(inp,res))},
                                    title=f"GUNITE — {inp.name}")


def figure_to_png_bytes(fig, dpi: int = 180) -> bytes:
    output = BytesIO()
    fig.savefig(output, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    output.seek(0)
    return output.getvalue()


def card(label: str, value: str):
    st.markdown(
        f'<div class="card"><div class="label">{label}</div><div class="value">{value}</div></div>',
        unsafe_allow_html=True,
    )


def calculate_saved_element(record):
    inp = GuniteInput(
        name=record.get("name", "K1"),
        x1b=float(record.get("x1b", 45)), x2b=float(record.get("x2b", 45)),
        y1b=float(record.get("y1b", 35)), y2b=float(record.get("y2b", 35)),
        fd_mm=float(record.get("fd_mm", 25)),
        corner_fd_mm=record.get("corner_fd_mm"), x_side_fd_mm=record.get("x_side_fd_mm"), y_side_fd_mm=record.get("y_side_fd_mm"),
        fs_mm=float(record.get("fs_mm", 10)), spacing_mm=int(record.get("spacing_mm", 100)),
        a1_cm=float(record.get("a1_cm", 1)), a2_cm=float(record.get("a2_cm", 0.5)),
        tgun_cm=float(record.get("tgun_cm", 8)), dh_cm=float(record.get("dh_cm", 5)),
        floor_height_m=float(record.get("floor_height_m", 3.46)), waiting_m=float(record.get("waiting_m", 1)),
        nx=int(record.get("nx", 3)), ny=int(record.get("ny", 4)),
        overlap_factor=float(record.get("overlap_factor", 80)), beam=None,
        development_option=str(record.get("development_option", "1a") or "1a"),
    )
    res = calculate(inp)
    return {"input": inp, "result": res, "order_rows": order_rows(inp, res),
            "floor_name": record.get("floor_name", "Στάθμη 1"),
            "study_dimensions": record.get("study_dimensions", ""),
            "remarks": record.get("remarks", ""),
            "designer_remarks": record.get("designer_remarks", ""),
            "length_cm": float(record.get("length_cm", (inp.x1b + inp.x2b))),
            "width_cm": float(record.get("width_cm", (inp.y1b + inp.y2b))),}


def save_current_project():
    project_id = database.save_project(
        st.session_state.project_name, st.session_state.project_elements,
        st.session_state.get("project_id"),
        floors=st.session_state.floors,
    )
    st.session_state.project_id = project_id
    st.session_state.project_saved = True


def load_project_into_session(project_id):
    data = database.load_project(int(project_id))
    if not data:
        return False
    st.session_state.project_id = data["id"]
    st.session_state.project_name = data["name"]
    st.session_state.project_elements = [calculate_saved_element(r) for r in data["elements"]]
    stored_floors = data.get("floors") or []
    if stored_floors:
        st.session_state.floors = stored_floors
    else:
        floor_names = []
        floors = []
        for item in st.session_state.project_elements:
            name = item.get("floor_name", "Στάθμη 1")
            if name not in floor_names:
                floor_names.append(name)
                inp = item["input"]
                floors.append({"name": name, "height": inp.floor_height_m, "tgun": inp.tgun_cm, "a1": inp.a1_cm, "a2": inp.a2_cm})
        st.session_state.floors = floors or [{"name": "Στάθμη 1", "height": 3.46, "tgun": 8.0, "a1": 1.0, "a2": 0.5}]
    st.session_state.selected_floor = 0
    st.session_state.current_step = 3 if st.session_state.project_elements else 0
    st.session_state.project_saved = True
    return True


def new_project():
    st.session_state.project_id = None
    st.session_state.project_name = "Νέο έργο"
    st.session_state.project_elements = []
    st.session_state.project_saved = False
    st.session_state.floors = [{"name": "Στάθμη 1", "height": 3.46, "tgun": 7.5, "a1": 1.0, "a2": 0.5}]
    st.session_state.selected_floor = 0
    st.session_state.edit_index = None
    reset_element()


def reset_element():
    keys = [
        "name", "study_dimensions", "remarks", "designer_remarks", "case_short", "development_option", "length_cm", "width_cm", "x1b", "x2b", "y1b", "y2b", "fd", "corner_fd", "x_side_fd", "y_side_fd", "fs", "spacing", "a1", "a2", "tgun", "dh",
        "nx", "ny", "floor_height", "waiting", "overlap_factor", "beam_enabled", "beam_side",
        "interrupted_bars", "beam_bottom_m", "beam_height_m", "lower_piece_m", "upper_piece_m",
    ]
    for key in keys:
        st.session_state.pop(key, None)
    st.session_state.current_step = 0


def next_element_name():
    existing = {str(e["input"].name).strip().upper() for e in st.session_state.get("project_elements", [])}
    n = 1
    while f"K{n}" in existing:
        n += 1
    return f"K{n}"


def prepare_new_element():
    reset_element()
    apply_selected_floor_defaults()
    st.session_state.name = next_element_name()
    # 4Π source example: the existing section is entered as TOTAL X and Y
    # (e.g. 45 × 35 cm), while the DOCX uses X1b=X2b=45 and Y1b=Y2b=35.
    st.session_state.length_cm = 45.0
    st.session_state.width_cm = 35.0
    st.session_state.x1b = 45.0
    st.session_state.x2b = 45.0
    st.session_state.y1b = 35.0
    st.session_state.y2b = 35.0
    st.session_state.study_dimensions = ""
    st.session_state.remarks = ""
    st.session_state.designer_remarks = ""
    st.session_state.jacket_case = "4πλευρος μανδύας"
    st.session_state.development_option = "1a"


def apply_selected_floor_defaults():
    floor = st.session_state.floors[st.session_state.selected_floor]
    st.session_state.floor_height = floor["height"]
    st.session_state.tgun = floor["tgun"]
    st.session_state.a1 = floor["a1"]
    st.session_state.a2 = floor["a2"]


def geometry_values(case: str):
    # For 4Π the DOCX uses TOTAL existing dimensions X and Y and stores
    # X1b=X2b=X, Y1b=Y2b=Y.
    if case == "4πλευρος μανδύας":
        total_x = float(st.session_state.get("length_cm", 45.0))
        total_y = float(st.session_state.get("width_cm", 35.0))
        # DOCX convention for 4Π: X1b=X2b=the full X dimension and
        # Y1b=Y2b=the full Y dimension. The calculator adds tgun/d ONCE
        # to each reported side dimension, giving +2*tgun overall.
        return total_x, total_x, total_y, total_y

    # 3Π/2Π remain unchanged for now; they are outside the current correction scope.
    length = float(st.session_state.get("length_cm", 90.0))
    width = float(st.session_state.get("width_cm", 70.0))
    if case == "3πλευρος μανδύας - ελεύθερη κάτω πλευρά":
        return length / 2.0, length / 2.0, 0.0, width
    return 0.0, length, 0.0, width


def build_input(case: str) -> GuniteInput:
    x1b, x2b, y1b, y2b = geometry_values(case)
    # Level is authoritative for the common vertical/jacketing parameters.
    floor = st.session_state.floors[st.session_state.selected_floor]
    # Beam/repair rules are intentionally not fabricated here. The supplied
    # theory contains no complete numerical rule for them.
    return GuniteInput(
        name=st.session_state.get("name", "K1"),
        x1b=x1b, x2b=x2b, y1b=y1b, y2b=y2b,
        fd_mm=float(st.session_state.get("fd", 25.0)),
        corner_fd_mm=float(st.session_state.get("corner_fd", st.session_state.get("fd", 25.0))),
        x_side_fd_mm=float(st.session_state.get("x_side_fd", st.session_state.get("fd", 25.0))),
        y_side_fd_mm=float(st.session_state.get("y_side_fd", st.session_state.get("fd", 25.0))),
        fs_mm=float(st.session_state.get("fs", 10.0)),
        spacing_mm=int(st.session_state.get("spacing", 100)),
        a1_cm=float(floor["a1"]),
        a2_cm=float(floor["a2"]),
        tgun_cm=float(floor["tgun"]),
        dh_cm=float(st.session_state.get("dh", 5.0)),
        floor_height_m=float(floor["height"]),
        waiting_m=float(st.session_state.get("waiting", 1.0)),
        nx=int(st.session_state.get("nx", 3)),
        ny=int(st.session_state.get("ny", 4)),
        overlap_factor=float(st.session_state.get("overlap_factor", 80.0)),
        beam=None,
        development_option=str(st.session_state.get("development_option", "1a") or "1a"),
    )


def render_live_preview(case: str):
    """Calculate the same object used by the final step and show it live."""
    try:
        inp = build_input(case)
        res = calculate(inp)
        col_a, col_b = st.columns(2)
        with col_a:
            card("Υφιστάμενη διατομή", f"{res.existing_width_m*100:.3f} × {res.existing_height_m*100:.3f} cm")
        with col_b:
            card("Τελική διατομή gunite", f"{res.gunite_width_m*100:.3f} × {res.gunite_height_m*100:.3f} cm")
        fig = drawing.draw_column_section(inp, res, size="Μεγάλο")
        st.pyplot(fig, width="stretch")
        import matplotlib.pyplot as plt
        plt.close(fig)
        st.caption("Η προεπισκόπηση χρησιμοποιεί ακριβώς τον ίδιο υπολογιστικό πυρήνα με το τελικό αποτέλεσμα.")
        return res
    except Exception as exc:
        st.warning(f"Η προεπισκόπηση δεν μπορεί ακόμη να υπολογιστεί: {exc}")
        return None


# ---------- Header / state ----------
if "project_elements" not in st.session_state:
    st.session_state.project_elements = []
if "current_step" not in st.session_state:
    st.session_state.current_step = 0
if "jacket_case" not in st.session_state:
    st.session_state.jacket_case = "4πλευρος μανδύας"
if "project_name" not in st.session_state:
    st.session_state.project_name = "Νέο έργο"
if "project_id" not in st.session_state:
    st.session_state.project_id = None
if "project_saved" not in st.session_state:
    st.session_state.project_saved = False
if "floors" not in st.session_state:
    st.session_state.floors = [{"name": "Στάθμη 1", "height": 3.46, "tgun": 7.5, "a1": 1.0, "a2": 0.5}]
if "selected_floor" not in st.session_state:
    st.session_state.selected_floor = 0
if "edit_index" not in st.session_state:
    st.session_state.edit_index = None
database.init_db()

st.markdown(
    '<div class="hero"><div class="hero-brand"><div class="hero-mark">G</div><div>'
    '<div class="section-kicker">ENGINEERING CALCULATOR</div>'
    '<h1>GUNITE</h1><div class="subtitle">Υπολογισμός μανδύα · Οπλισμός · Προμέτρηση · Παραγγελία</div>'
    '</div></div>'
    '<div class="context"><span class="chip">Έργο → Στάθμη → Κολώνα</span><span class="chip">4Π έλεγχος βάσει GUNITE_Logic</span><span class="chip">Excel & PDF</span></div></div>',
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("Έργο")
    st.caption("Οργάνωση έργου ανά στάθμη → στοιχεία. Για τον 4Π, οι κοινές παράμετροι της στάθμης είναι οι ενεργές τιμές του υπολογισμού.")
    project_options = database.list_projects()
    option_map = {f"{p['name']}  ·  #{p['id']}": p['id'] for p in project_options}
    labels = ["— Επιλογή αποθηκευμένου έργου —"] + list(option_map.keys())
    current_label = next((k for k, v in option_map.items() if v == st.session_state.get("project_id")), labels[0])
    selected_label = st.selectbox("Άνοιγμα έργου", labels, index=labels.index(current_label) if current_label in labels else 0)
    if selected_label != labels[0] and option_map.get(selected_label) != st.session_state.get("project_id"):
        if st.button("📂 Άνοιγμα", use_container_width=True):
            load_project_into_session(option_map[selected_label]); st.rerun()
    st.text_input("Όνομα έργου", key="project_name")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("💾 Αποθήκευση", use_container_width=True): save_current_project(); st.rerun()
    with c2:
        if st.button("＋ Νέο", use_container_width=True): new_project(); st.rerun()
    if st.session_state.get("project_saved"): st.success("Αποθηκευμένο")
    st.metric("Στοιχεία στο έργο", len(st.session_state.project_elements))

    st.divider()
    st.subheader("🌳 Στάθμες")
    st.caption("Η Στάθμη = ο όροφος/επίπεδο του κτιρίου. Εδώ βάζεις μόνο τις κοινές τιμές που ισχύουν για τις κολώνες αυτού του ορόφου. Οι κολώνες έχουν ξεχωριστά τις δικές τους διαστάσεις και οπλισμούς.")
    for fi, floor in enumerate(st.session_state.floors):
        floor_title = f"{fi + 1}. {floor['name']}"
        with st.expander(floor_title, expanded=(fi == st.session_state.selected_floor)):
            st.markdown("**🏢 Κοινά δεδομένα του ορόφου**")
            old_floor_name = floor["name"]
            floor["name"] = st.text_input("Όνομα στάθμης", value=floor["name"], key=f"floor_name_{fi}")
            if floor["name"] != old_floor_name:
                for _element in st.session_state.project_elements:
                    if _element.get("floor_name") == old_floor_name:
                        _element["floor_name"] = floor["name"]
                st.session_state.project_saved = False
            floor["height"] = st.number_input("Ύψος ορόφου / στάθμης (m)", min_value=0.01, value=float(floor["height"]), step=0.01, format="%.3f", key=f"floor_h_{fi}")
            floor["tgun"] = st.number_input("Πάχος Gunite tgun (cm)", min_value=0.0, value=float(floor["tgun"]), step=0.5, format="%.3f", key=f"floor_t_{fi}")
            floor["a1"] = st.number_input("Απόσταση a1 (cm)", min_value=0.0, value=float(floor["a1"]), step=0.1, format="%.3f", key=f"floor_a1_{fi}")
            floor["a2"] = st.number_input("Απόσταση a2 (cm)", min_value=0.0, value=float(floor["a2"]), step=0.1, format="%.3f", key=f"floor_a2_{fi}")
            if st.button("Ενεργή στάθμη", key=f"select_floor_{fi}", use_container_width=True):
                st.session_state.selected_floor = fi
                reset_element()
                apply_selected_floor_defaults()
                st.rerun()

            st.caption("Ύψος, tgun, a1 και a2 είναι κοινά δεδομένα του ορόφου. Κάθε νέα κολώνα της στάθμης τα χρησιμοποιεί αυτόματα.")

            floor_items = [(i, e) for i, e in enumerate(st.session_state.project_elements) if e.get("floor_name") == floor["name"]]
            st.markdown("**Στοιχεία στάθμης**")
            if not floor_items:
                st.caption("— κανένα στοιχείο —")
            else:
                for local_pos, (global_i, e) in enumerate(floor_items, start=1):
                    inp_e = e["input"]
                    res_e = e["result"]
                    case_short = "4Π" if res_e.case.startswith("4") else "3Π" if "3πλευρος" in res_e.case else "2Π"
                    label = f"K{inp_e.name[1:] if inp_e.name.upper().startswith('K') and inp_e.name[1:].isdigit() else inp_e.name}  ·  {e.get('length_cm', inp_e.x1b+inp_e.x2b):.3f}×{e.get('width_cm', inp_e.y1b+inp_e.y2b):.3f}  ·  {case_short}"
                    cedit, cdel = st.columns([5, 1])
                    with cedit:
                        if st.button(label, key=f"tree_el_{fi}_{global_i}", use_container_width=True, help="Άνοιγμα στοιχείου για επεξεργασία"):
                            st.session_state.selected_floor = fi
                            st.session_state.edit_index = global_i
                            st.session_state.jacket_case = res_e.case
                            st.session_state.development_option = getattr(inp_e, "development_option", "1a") or "1a"
                            st.session_state.name = inp_e.name
                            st.session_state.study_dimensions = e.get("study_dimensions", "")
                            st.session_state.remarks = e.get("remarks", "")
                            st.session_state.designer_remarks = e.get("designer_remarks", "")
                            st.session_state.length_cm = e.get("length_cm", inp_e.x1b + inp_e.x2b)
                            st.session_state.width_cm = e.get("width_cm", inp_e.y1b + inp_e.y2b)
                            st.session_state.x1b = inp_e.x1b
                            st.session_state.x2b = inp_e.x2b
                            st.session_state.y1b = inp_e.y1b
                            st.session_state.y2b = inp_e.y2b
                            for k, v in {"fd": inp_e.fd_mm, "corner_fd": inp_e.corner_fd_mm or inp_e.fd_mm, "x_side_fd": inp_e.x_side_fd_mm or inp_e.fd_mm, "y_side_fd": inp_e.y_side_fd_mm or inp_e.fd_mm, "fs": inp_e.fs_mm, "spacing": inp_e.spacing_mm, "a1": inp_e.a1_cm, "a2": inp_e.a2_cm, "tgun": inp_e.tgun_cm, "dh": inp_e.dh_cm, "nx": inp_e.nx, "ny": inp_e.ny, "floor_height": inp_e.floor_height_m, "waiting": inp_e.waiting_m, "overlap_factor": inp_e.overlap_factor}.items():
                                st.session_state[k] = v
                            st.session_state.current_step = 1
                            st.rerun()
                    with cdel:
                        if st.button("🗑️", key=f"del_el_{fi}_{global_i}", help="Αφαίρεση στοιχείου"):
                            st.session_state.project_elements.pop(global_i)
                            st.session_state.project_saved = False
                            st.session_state.edit_index = None
                            save_current_project()
                            st.rerun()
            if st.button("＋ Προσθήκη στοιχείου", key=f"add_el_floor_{fi}", use_container_width=True):
                st.session_state.selected_floor = fi
                st.session_state.edit_index = None
                prepare_new_element()
                st.session_state.current_step = 1
                st.rerun()

    if st.button("＋ Προσθήκη στάθμης", use_container_width=True):
        n = len(st.session_state.floors) + 1
        st.session_state.floors.append({"name": f"Στάθμη {n}", "height": 3.46, "tgun": 7.5, "a1": 1.0, "a2": 0.5})
        st.session_state.selected_floor = n - 1
        reset_element()
        apply_selected_floor_defaults()
        st.rerun()

    if st.session_state.get("project_id") and st.button("🗑️ Διαγραφή αποθηκευμένου έργου", use_container_width=True):
        database.delete_project(st.session_state.project_id); new_project(); st.rerun()
    if database.DB_PATH.exists():
        st.download_button("🗄️ Αντίγραφο βάσης (.db)", data=database.DB_PATH.read_bytes(), file_name="gunite_projects_backup.db", mime="application/x-sqlite3", use_container_width=True)
    if st.session_state.get("project_id"):
        try:
            export_path = database.DB_PATH.parent / f"{st.session_state.project_name or 'project'}.gunite"
            project_io.export_project(st.session_state.project_id, export_path)
            st.download_button("📦 Εξαγωγή τρέχοντος έργου", data=export_path.read_bytes(), file_name=export_path.name, mime="application/json", use_container_width=True)
        except Exception as exc: st.caption(f"Δεν έγινε εξαγωγή έργου: {exc}")
    uploaded_project = st.file_uploader("📦 Εισαγωγή έργου (.gunite)", type=["gunite"], key="project_import")
    if uploaded_project is not None and st.button("Εισαγωγή αποθηκευμένου έργου", use_container_width=True):
        import tempfile
        try:
            with tempfile.NamedTemporaryFile(suffix=".gunite", delete=False) as tmp:
                tmp.write(uploaded_project.getvalue()); tmp_path=Path(tmp.name)
            imported_id=project_io.import_project(tmp_path); tmp_path.unlink(missing_ok=True); load_project_into_session(imported_id); st.rerun()
        except Exception as exc: st.error(f"Η εισαγωγή απέτυχε: {exc}")
    if st.button("🧹 Καθαρισμός τρέχοντος έργου", use_container_width=True): st.session_state.project_elements=[]; st.session_state.project_saved=False; reset_element(); st.rerun()

steps = ["01 · Στάθμη & στοιχεία", "02 · Δεδομένα κολώνας", "03 · Αποτελέσματα", "04 · Παραγγελία"]
current = st.session_state.current_step
workflow_html = '<div class="workflow">'
for i, label in enumerate(steps):
    cls = "workflow-item active" if i == current else "workflow-item"
    num, name = label.split(" · ", 1)
    workflow_html += f'<div class="{cls}"><div class="num">{num}</div><div class="name">{name}</div></div>'
    if i < len(steps)-1:
        workflow_html += '<div class="workflow-arrow">›</div>'
workflow_html += '</div>'
st.markdown(workflow_html, unsafe_allow_html=True)
step = st.radio("Βήμα εργασίας", steps, index=current, horizontal=True, label_visibility="collapsed")
st.session_state.current_step = steps.index(step)
st.progress((st.session_state.current_step + 1) / len(steps), text=f"Βήμα {st.session_state.current_step + 1} από 4")

# ---------- Step 1 ----------
if st.session_state.current_step == 0:
    st.subheader("1. Όροφος / Στάθμη → Κολώνες")
    st.markdown(
        "**Πώς λειτουργεί:** η Στάθμη είναι ο όροφος. Πρώτα ορίζεις τις κοινές παραμέτρους του ορόφου και μετά δημιουργείς τις κολώνες που ανήκουν σε αυτόν. Οι διαστάσεις και ο οπλισμός κάθε κολώνας μπαίνουν ξεχωριστά.")

    floor = st.session_state.floors[st.session_state.selected_floor]
    st.markdown(f"### 📍 Ενεργή στάθμη: {floor['name']}")
    st.info(
        f"**Κοινά δεδομένα που χρησιμοποιεί κάθε νέο στοιχείο αυτής της στάθμης:** "
        f"Ύψος = {floor['height']:.3f} m · tgun = {floor['tgun']:.3f} cm · "
        f"a1 = {floor['a1']:.3f} cm · a2 = {floor['a2']:.3f} cm. "
        "Τα δεδομένα αυτά είναι κοινά της στάθμης· δεν είναι οι διαστάσεις της κολώνας.")

    st.markdown("#### Τι συμπληρώνεις πού")
    guide = pd.DataFrame([
        {"Βήμα": "Στάθμη", "Συμπληρώνεις": "Ύψος, tgun, a1, a2", "Χρήση": "Ενεργές κοινές παράμετροι για τον υπολογισμό"},
        {"Βήμα": "Κολώνα", "Συμπληρώνεις": "X, Y και οπλισμό της συγκεκριμένης κολώνας", "Χρήση": "Υφιστάμενη διατομή· για 4Π: X1b=X2b=X και Y1b=Y2b=Y"},
        {"Βήμα": "Οπλισμός", "Συμπληρώνεις": "Φd, Φs, s, nx, ny", "Χρήση": "Υπολογισμός συνδετήρων και διαμήκων"},
        {"Βήμα": "Αποτέλεσμα", "Υπολογίζονται": "d, Xg/Yg, Xs/Ys, c, ανάπτυγμα, εμπορικός μανδύας", "Χρήση": "Έλεγχος πριν την παραγγελία"},
    ])
    st.dataframe(guide, hide_index=True, width="stretch")
    st.caption("Για τον 4Π ακολουθούμε το παράδειγμα του DOCX: εδώ δίνεις την πραγματική συνολική διατομή της υφιστάμενης κολώνας (π.χ. X=45 cm, Y=35 cm). Η εφαρμογή κρατά X1b=X2b=45 και Y1b=Y2b=35, όπως στο παράδειγμα της πηγής — όχι ως μισά.")

    rows = []
    floor_items = [(i, e) for i, e in enumerate(st.session_state.project_elements) if e.get("floor_name") == floor["name"]]
    for local_pos, (_, e) in enumerate(floor_items, start=1):
        rows.append(element_schedule_row(e, local_pos))

    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch", height=360)
        st.markdown("**Άνοιγμα στοιχείου για επεξεργασία**")
        for local_pos, (global_i, e) in enumerate(floor_items, start=1):
            inp_e = e["input"]
            if st.button(f"✎  {inp_e.name} — επεξεργασία", key=f"main_edit_{global_i}", use_container_width=True):
                st.session_state.edit_index = global_i
                st.session_state.jacket_case = e["result"].case
                st.session_state.development_option = getattr(inp_e, "development_option", "1a") or "1a"
                st.session_state.name = inp_e.name
                st.session_state.study_dimensions = e.get("study_dimensions", "")
                st.session_state.remarks = e.get("remarks", "")
                st.session_state.designer_remarks = e.get("designer_remarks", "")
                st.session_state.length_cm = e.get("length_cm", inp_e.x1b + inp_e.x2b)
                st.session_state.width_cm = e.get("width_cm", inp_e.y1b + inp_e.y2b)
                st.session_state.x1b = inp_e.x1b
                st.session_state.x2b = inp_e.x2b
                st.session_state.y1b = inp_e.y1b
                st.session_state.y2b = inp_e.y2b
                for k, v in {"fd": inp_e.fd_mm, "corner_fd": inp_e.corner_fd_mm or inp_e.fd_mm, "x_side_fd": inp_e.x_side_fd_mm or inp_e.fd_mm, "y_side_fd": inp_e.y_side_fd_mm or inp_e.fd_mm, "fs": inp_e.fs_mm, "spacing": inp_e.spacing_mm, "a1": inp_e.a1_cm, "a2": inp_e.a2_cm, "tgun": inp_e.tgun_cm, "dh": inp_e.dh_cm, "nx": inp_e.nx, "ny": inp_e.ny, "floor_height": inp_e.floor_height_m, "waiting": inp_e.waiting_m, "overlap_factor": inp_e.overlap_factor}.items():
                    st.session_state[k] = v
                st.session_state.current_step = 1
                st.rerun()
    else:
        st.info("Η στάθμη δεν έχει ακόμη στοιχεία.")

    if st.button("＋ Προσθήκη στοιχείου στη στάθμη", type="primary", use_container_width=True):
        st.session_state.edit_index = None
        prepare_new_element()
        st.session_state.current_step = 1
        st.rerun()

# ---------- Step 2 ----------
elif st.session_state.current_step == 1:
    case = st.session_state.jacket_case
    st.markdown('<div class="section-kicker">02 · INPUTS</div><div class="section-title">Δεδομένα κολώνας</div><div class="section-help">Συμπλήρωσε μόνο τα στοιχεία που αφορούν τη συγκεκριμένη κολώνα. Οι κοινές παράμετροι της ενεργής στάθμης εμφανίζονται κλειδωμένες ως αναφορά.</div>', unsafe_allow_html=True)
    left, right = st.columns([1.0, 1.25], gap="large")

    with left:
        with st.expander("Α. Στοιχείο & διατομή", expanded=True):
            st.text_input("Στοιχείο", key="name", help="π.χ. K1, K2, K3 — για νέο στοιχείο δίνεται αυτόματα το επόμενο K.")
            st.text_input("Διαστάσεις Μελέτης", key="study_dimensions", placeholder="π.χ. 35/35 ή 40/40", help="Ακριβώς όπως στον πίνακα μελέτης. Δεν χρησιμοποιείται αυθαίρετα στον υπολογισμό.")
            c_rem1, c_rem2 = st.columns(2)
            c_rem1.text_input("Παρατηρήσεις", key="remarks", placeholder="")
            c_rem2.text_input("Παρατηρήσεις μελετητή", key="designer_remarks", placeholder="")
            if case == "4πλευρος μανδύας":
                st.markdown("**Υφιστάμενη διατομή — 4Π**")
                st.caption("Σύμφωνα με το αριθμητικό παράδειγμα της πηγής, εδώ γράφεις τις ΣΥΝΟΛΙΚΕΣ διαστάσεις της παλιάς κολώνας: X και Y. Δεν γράφεις μισές διαστάσεις.")
                c1, c2 = st.columns(2)
                c1.number_input("Συνολικό X υφιστάμενης (cm)", min_value=0.001, step=0.5, value=float(st.session_state.get("length_cm", 45.0)), format="%.3f", key="length_cm")
                c2.number_input("Συνολικό Y υφιστάμενης (cm)", min_value=0.001, step=0.5, value=float(st.session_state.get("width_cm", 35.0)), format="%.3f", key="width_cm")
                st.success(f"Υφιστάμενη διατομή: **{float(st.session_state.get('length_cm',45)):.3f} × {float(st.session_state.get('width_cm',35)):.3f} cm** → για 4Π χρησιμοποιούνται X1b=X2b={float(st.session_state.get('length_cm',45)):.3f} cm και Y1b=Y2b={float(st.session_state.get('width_cm',35)):.3f} cm, όπως στο DOCX.")
            else:
                st.markdown("**Πραγματική διατομή της υφιστάμενης κολώνας**")
                st.caption("Για 3Π/2Π χρησιμοποιούνται οι διαστάσεις της αντίστοιχης γεωμετρίας.")
                c1, c2 = st.columns(2)
                c1.number_input("Συνολικό μήκος X (cm)", min_value=0.001, step=0.001, value=float(st.session_state.get("length_cm", 90.0)), format="%.3f", key="length_cm")
                c2.number_input("Συνολικό πλάτος Y (cm)", min_value=0.001, step=0.001, value=float(st.session_state.get("width_cm", 70.0)), format="%.3f", key="width_cm")
            case_options = {
                "4Π": "4πλευρος μανδύας",
                "3Π": "3πλευρος μανδύας - ελεύθερη κάτω πλευρά",
                "2Π": "γωνιακός μανδύας",
            }
            current_short = "4Π" if case.startswith("4") else "3Π" if "3πλευρος" in case else "2Π"
            selected_short = st.selectbox("Πλευρές μανδύα", list(case_options.keys()), index=list(case_options.keys()).index(current_short), key="case_short")
            st.session_state.jacket_case = case_options[selected_short]
            case = st.session_state.jacket_case
            if case == "4πλευρος μανδύας":
                options = {
                    "1a — Π με οριζόντια πλευρά X": "1a",
                    "1b — Π με οριζόντια πλευρά Y": "1b",
                }
                current_option = st.session_state.get("development_option", "1a")
                if current_option not in options.values():
                    current_option = "1a"
                selected_option = st.selectbox(
                    "Τρόπος ανάπτυξης 4Π", list(options.keys()),
                    index=list(options.values()).index(current_option),
                    key="development_option_label"
                )
                st.session_state.development_option = options[selected_option]
                st.caption("1a: 55 + 22,5 + 22,5 = 100 cm. 1b: 45 + 27,5 + 27,5 = 100 cm. Στη συνέχεια προστίθεται η απαίτηση συγκόλλησης και επιλέγεται το εμπορικό ανάπτυγμα.")
            if case == "3πλευρος μανδύας - ελεύθερη κάτω πλευρά":
                st.caption("3Π: η κάτω πλευρά είναι ελεύθερη.")
            elif case == "γωνιακός μανδύας":
                st.caption("2Π: γωνιακός μανδύας.")

        with st.expander("Β. Οπλισμός & gunite", expanded=True):
            diameters_d = [8,10,12,14,16,18,20,22,25,28,32]
            fd_default = int(st.session_state.get("fd", 25))
            fd = st.selectbox("Φd — διαμήκης ράβδος (mm)", diameters_d, index=diameters_d.index(fd_default) if fd_default in diameters_d else 0, key="fd")
            diameters = available_diameters()
            fs_default = int(st.session_state.get("fs", 10))
            fs = st.selectbox("Φs — συνδετήρας / μανδύας (mm)", diameters, index=diameters.index(fs_default) if fs_default in diameters else 0, key="fs")
            spacings = available_spacings(int(fs))
            spacing_default = int(st.session_state.get("spacing", spacings[0])) if spacings else 100
            spacing = st.selectbox("Απόσταση s από τον πραγματικό κατάλογο (mm)", spacings or [100], index=(spacings.index(spacing_default) if spacing_default in spacings else 0), key="spacing")
            floor_now = st.session_state.floors[st.session_state.selected_floor]
            c1, c2 = st.columns(2)
            c1.metric("a1 — από Στάθμη (cm)", f"{floor_now['a1']:.3f}")
            c2.metric("a2 — από Στάθμη (cm)", f"{floor_now['a2']:.3f}")
            c1, c2 = st.columns(2)
            c1.metric("Πάχος Gunite — από Στάθμη (cm)", f"{floor_now['tgun']:.3f}")
            c2.metric("Ύψος στάθμης — από Στάθμη (m)", f"{floor_now['height']:.3f}")
            st.caption("Οι παραπάνω 4 τιμές δεν συμπληρώνονται στην κολώνα. Τις αλλάζεις μόνο από την ενεργή Στάθμη.")
            if case == "4πλευρος μανδύας":
                st.number_input("dh (Hilti) (cm)", min_value=0.0, step=0.5, value=float(st.session_state.get("dh", 5.0)), format="%.3f", key="dh", disabled=True, help="Δεν συμμετέχει στους υπολογισμούς 4Π σύμφωνα με την παρατήρηση.")
                st.caption("dh Hilti: ανενεργό για 4Π. Θα ενεργοποιείται μόνο σε 3Π/2Π.")
            else:
                st.number_input("dh (Hilti) (cm)", min_value=0.0, step=0.5, value=float(st.session_state.get("dh", 5.0)), format="%.3f", key="dh")
            st.caption(f"Πηγή καταλόγου: {catalog.CATALOG_SOURCE}")

        with st.expander("Γ. Διαμήκεις ράβδοι", expanded=True):
            c1, c2 = st.columns(2)
            c1.number_input("nx — ράβδοι κατά X", min_value=1, step=1, value=int(st.session_state.get("nx", 3)), key="nx")
            c2.number_input("ny — ράβδοι κατά Y", min_value=1, step=1, value=int(st.session_state.get("ny", 4)), key="ny")
            st.markdown("**Διαφορετικές διατομές ανά θέση**")
            st.caption("Παράδειγμα: 4Φ20 στις γωνίες + Φ16 στις ενδιάμεσες πλευρές. Οι ποσότητες υπολογίζονται από nx/ny και την περίπτωση του μανδύα.")
            c1, c2, c3 = st.columns(3)
            for col, label, key in [(c1,"Γωνίες Φ (mm)","corner_fd"),(c2,"Πλευρές X Φ (mm)","x_side_fd"),(c3,"Πλευρές Y Φ (mm)","y_side_fd")]:
                default=int(st.session_state.get(key, fd))
                col.selectbox(label, diameters_d, index=diameters_d.index(default) if default in diameters_d else diameters_d.index(int(fd)), key=key)
            st.info("Η εφαρμογή θα εμφανίσει αμέσως στο αποτέλεσμα πόσες ράβδοι ανήκουν σε κάθε ομάδα.")
            st.number_input("Αναμονή (m)", min_value=0.0, step=0.05, value=float(st.session_state.get("waiting", 1.0)), format="%.3f", key="waiting")
            st.number_input("Συντελεστής υπερκάλυψης ×Φd", min_value=0.0, step=1.0, value=float(st.session_state.get("overlap_factor", 80.0)), key="overlap_factor")

        with st.expander("Δ. Δοκός / εμπόδιο", expanded=False):
            st.info("ON HOLD — το αφήνουμε ανενεργό μέχρι να οριστεί από τις πηγές ο πλήρης κανόνας για θέση δοκού, διακοπή και μαντάρισμα ράβδων. Δεν γίνεται κανένας αυθαίρετος υπολογισμός.")

    with right:
        st.markdown("### Ζωντανό σχέδιο")
        render_live_preview(case)

    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        if st.button("← Πίσω", use_container_width=True):
            st.session_state.current_step = 0
            st.rerun()
    with c2:
        if st.button("Έλεγχος αποτελέσματος →", type="primary", use_container_width=True):
            st.session_state.current_step = 2
            st.rerun()

# ---------- Step 3 ----------
elif st.session_state.current_step == 2:
    st.markdown('<div class="section-kicker">03 · CHECK</div><div class="section-title">Αποτελέσματα & τεχνικός έλεγχος</div><div class="section-help">Εδώ ελέγχεις τις διαστάσεις που προέκυψαν, το σχέδιο, τις διαμήκεις ράβδους και τα τεμάχια του καταλόγου.</div>', unsafe_allow_html=True)
    try:
        inp = build_input(st.session_state.jacket_case)
        res = calculate(inp)
        current_result_df = pd.DataFrame(result_table(inp, res))
        current_order_rows = order_rows(inp, res)

        if res.catalog_name is None or (res.p_development_bottom is not None and res.catalog_name_bottom is None):
            st.error("Δεν υπάρχει κατάλληλο πραγματικό προϊόν στον κατάλογο για ένα από τα απαιτούμενα τεμάχια. Δεν δημιουργείται πλασματικός κωδικός.")
        else:
            st.success("Ο υπολογισμός ολοκληρώθηκε και τα τεμάχια αντιστοιχίστηκαν σε πραγματικές εγγραφές του καταλόγου.")

        cols = st.columns(5)
        values = [
            ("Υφιστάμενη", f"{res.existing_width_m*100:.3f} × {res.existing_height_m*100:.3f} cm"),
            ("Τελική gunite", f"{res.gunite_width_m*100:.3f} × {res.gunite_height_m*100:.3f} cm"),
            ("Διαμήκεις", f"{res.longitudinal_bars_total} τεμ."),
            ("Gunite", f"{res.gunite_volume_m3:.3f} m³"),
            ("Μανδύες", f"{2 if res.p_development_bottom else 1} τεμ."),
        ]
        for col, (label, value) in zip(cols, values):
            with col:
                card(label, value)

        left, right = st.columns([1.25, 1], gap="large")
        with left:
            st.markdown("### Σχέδιο διατομής")
            fig = drawing.draw_column_section(inp, res, size="Μεγάλο")
            st.pyplot(fig, width="stretch")
            section_png = figure_to_png_bytes(fig)
            import matplotlib.pyplot as plt
            plt.close(fig)
            st.download_button(
                "📐 Λήψη διατομής PNG", section_png,
                file_name=f"{inp.name}_diatomi.png", mime="image/png",
                use_container_width=True, key=f"section_png_{inp.name}_{len(st.session_state.project_elements)}"
            )

            st.markdown("### Κατακόρυφη όψη διαμήκων")
            fig_elev = drawing.draw_vertical_reinforcement(inp, res, size="Μεγάλο")
            st.pyplot(fig_elev, width="stretch")
            elev_png = figure_to_png_bytes(fig_elev)
            plt.close(fig_elev)
            st.download_button(
                "📏 Λήψη κατακόρυφης όψης PNG", elev_png,
                file_name=f"{inp.name}_katakoryfi.png", mime="image/png",
                use_container_width=True, key=f"elev_png_{inp.name}_{len(st.session_state.project_elements)}"
            )

            st.markdown("### Τεχνικό φύλλο")
            tech_fig = drawing.draw_technical_sheet(inp, res, size="Μεγάλο")
            st.pyplot(tech_fig, width="stretch")
            tech_png = figure_to_png_bytes(tech_fig, dpi=220)
            plt.close(tech_fig)
            st.download_button(
                "📄 Λήψη τεχνικού φύλλου PNG", tech_png,
                file_name=f"{inp.name}_techniko_fyllo.png", mime="image/png",
                use_container_width=True, key=f"tech_png_{inp.name}_{len(st.session_state.project_elements)}"
            )
        with right:
            st.markdown("### Τεμάχια μανδύα")
            jacket_data = []
            specs = [(res.p_development.label, res.catalog_name, res.catalog_weight_kg)]
            if res.p_development_bottom is not None:
                specs.append((res.p_development_bottom.label, res.catalog_name_bottom, res.catalog_weight_kg_bottom))
            for position, name, weight in specs:
                dev_obj = res.p_development if position == res.p_development.label else res.p_development_bottom
                jacket_data.append({"Θέση": position, "Γεωμετρικό (cm)": round(dev_obj.geometric_required_m*100,1), "Με συγκόλληση (cm)": round(dev_obj.required_m*100,1), "Προϊόν": name or "—", "Τεμάχια": 1, "Υπόλοιπο/άκρο (cm)": round(dev_obj.surplus_each_end_cm,1), "Βάρος/τεμ. (kg)": weight or 0.0})
            st.dataframe(pd.DataFrame(jacket_data), width="stretch", hide_index=True)

            st.markdown("### Βασικοί υπολογισμοί")
            with st.expander("Διαστάσεις & ανάπτυγμα 4Π", expanded=True):
                if res.case == "4πλευρος μανδύας":
                    st.write(f"Τρόπος ανάπτυξης: **{getattr(inp, 'development_option', '1a')}**")
                    st.write(f"Π γεωμετρικό: **{res.p_development.geometric_required_m*100:.1f} cm** · + συγκόλληση **{2*res.p_development.welding_half_cm:.1f} cm** → απαιτούμενο **{res.p_development.required_m*100:.1f} cm**")
                    if res.catalog_name:
                        st.write(f"Εμπόριο: **{res.catalog_name}** · υπόλοιπο **{res.p_development.surplus_each_end_cm:.1f} cm/άκρο**")
                st.write(f"Υφιστάμενη: **{res.existing_width_m*100:.1f} × {res.existing_height_m*100:.1f} cm**")
                st.write(f"Gunite: **{res.gunite_width_m*100:.1f} × {res.gunite_height_m*100:.1f} cm**")
                st.write(f"d = **{res.d_cm:.2f} cm**")
            with st.expander("Πλήθος διαμήκων", expanded=False):
                st.write(f"Τύπος περίπτωσης: **{res.case}**")
                st.write(f"nx = {inp.nx}, ny = {inp.ny} → **{res.longitudinal_bars_total} ράβδοι**")
                group_rows = []
                for group, info in res.longitudinal_groups.items():
                    if int(info["count"]) > 0:
                        group_rows.append({"Ομάδα": group, "Τεμάχια": int(info["count"]), "Φ (mm)": f"{info['diameter_mm']:.0f}", "Μήκος/τεμ. (m)": f"{info['length_m']:.3f}", "Υπερκάλυψη (m)": f"{info['overlap_m']:.3f}"})
                st.dataframe(pd.DataFrame(group_rows), width="stretch", hide_index=True)
            with st.expander("Προμέτρηση", expanded=False):
                st.write(f"Περίμετρος gunite: **{res.gunite_perimeter_m:.2f} m**")
                st.write(f"Επιφάνεια gunite: **{res.gunite_surface_m2:.2f} m²**")
                st.write(f"Όγκος gunite: **{res.gunite_volume_m3:.3f} m³**")
            with st.expander("Δοκός / εμπόδιο", expanded=False):
                st.info("ON HOLD — δεν εκτελείται υπολογισμός δοκού/εμποδίου μέχρι να οριστούν οι απαιτούμενοι κανόνες από τις πηγές.")

        st.markdown(f"### Παραγγελία μόνο για το {inp.name}")
        st.caption("Αυτές οι γραμμές αφορούν μόνο αυτό το στοιχείο. Όταν το προσθέσεις στο έργο, θα μεταφερθούν και στη συγκεντρωτική παραγγελία.")
        st.dataframe(pd.DataFrame(current_order_rows), width="stretch", hide_index=True)

        element_excel = build_element_excel(inp, res)
        element_pdf = create_project_pdf(
            pd.DataFrame([{"Στοιχείο": inp.name, "Περίπτωση": res.case,
                           "Διατομή gunite": f"{res.gunite_width_m*100:.3f} × {res.gunite_height_m*100:.3f} cm",
                           "Διαμήκεις": f"{res.longitudinal_bars_total} Φ{inp.fd_mm:.0f}",
                           "Gunite (m³)": res.gunite_volume_m3}]),
            pd.DataFrame(current_order_rows), jacket_dataframe(inp,res),
            input_df=input_dataframe(inp), calculation_df=calculation_dataframe(inp,res))
        cexp1, cexp2 = st.columns(2)
        with cexp1:
            st.download_button("📥 Κατέβασε αυτό το στοιχείο σε Excel", element_excel,
                               file_name=f"{inp.name}_GUNITE.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True)
        with cexp2:
            st.download_button("📄 Κατέβασε αυτό το στοιχείο σε PDF", element_pdf,
                               file_name=f"{inp.name}_GUNITE.pdf", mime="application/pdf",
                               use_container_width=True)

        st.divider()
        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("← Διόρθωση δεδομένων", use_container_width=True):
                st.session_state.current_step = 1
                st.rerun()
        with c2:
            action_label = "💾 Αποθήκευση αλλαγών στοιχείου" if st.session_state.edit_index is not None else "➕ Προσθήκη στο έργο"
            if st.button(action_label, type="primary", use_container_width=True):
                floor = st.session_state.floors[st.session_state.selected_floor]
                item = {"input": inp, "result": res, "order_rows": current_order_rows,
                        "floor_name": floor["name"], "study_dimensions": str(st.session_state.get("study_dimensions", "")).strip(),
                        "remarks": str(st.session_state.get("remarks", "")).strip(),
                        "designer_remarks": str(st.session_state.get("designer_remarks", "")).strip(),
                        "length_cm": float(inp.x1b if inp.x1b > 0 and abs(inp.x1b-inp.x2b) < 1e-9 else inp.x1b + inp.x2b), "width_cm": float(inp.y1b if inp.y1b > 0 and abs(inp.y1b-inp.y2b) < 1e-9 else inp.y1b + inp.y2b)}
                if st.session_state.edit_index is None:
                    st.session_state.project_elements.append(item)
                else:
                    st.session_state.project_elements[st.session_state.edit_index] = item
                    st.session_state.edit_index = None
                st.session_state.project_saved = False
                save_current_project()
                st.session_state.current_step = 3
                st.rerun()
        with c3:
            if st.button("🔄 Νέο στοιχείο", use_container_width=True):
                st.session_state.edit_index = None
                prepare_new_element()
                st.rerun()
    except Exception as exc:
        st.error(str(exc))
        if st.button("← Επιστροφή στα δεδομένα"):
            st.session_state.current_step = 1
            st.rerun()

# ---------- Step 4 ----------
else:
    st.markdown(f'<div class="section-kicker">04 · OUTPUT</div><div class="section-title">Παραγγελία έργου</div><div class="section-help">Συγκεντρωτική εικόνα όλων των στοιχείων, των τεμαχίων καταλόγου και των συνολικών ποσοτήτων.</div>', unsafe_allow_html=True)
    st.info("Εδώ βλέπεις όλα τα στοιχεία του έργου μαζί. Η ‘Παραγγελία στοιχείου’ αφορά μόνο το συγκεκριμένο Κ1/Κ2/... . Η ‘Συγκεντρωτική παραγγελία’ αθροίζει τα ίδια προϊόντα από όλα τα στοιχεία του έργου.")
    if not st.session_state.project_elements:
        st.info("Δεν υπάρχει ακόμη στοιχείο στο έργο.")
        if st.button("← Νέο στοιχείο", type="primary"):
            st.session_state.current_step = 0
            st.rerun()
    else:
        elements_summary = []
        all_order_rows = []
        for index, item in enumerate(st.session_state.project_elements, start=1):
            all_order_rows.extend(item["order_rows"])
            row = element_schedule_row(item, index)
            row["Στάθμη"] = item.get("floor_name", "—")
            row["Gunite (m³)"] = item["result"].gunite_volume_m3
            elements_summary.append(row)
        elements_df = pd.DataFrame(elements_summary)
        project_order_df = pd.DataFrame(all_order_rows)
        if "Σύνολο kg" not in project_order_df:
            project_order_df["Σύνολο kg"] = 0.0
        project_order_df["Σύνολο kg"] = project_order_df["Σύνολο kg"].fillna(0.0)

        summary_df = (
            project_order_df.groupby(["Τύπος", "Περιγραφή", "Μήκος (m)"], as_index=False)
            .agg({"Τεμάχια": "sum", "Σύνολο kg": "sum"})
            .sort_values(["Τύπος", "Περιγραφή", "Μήκος (m)"], na_position="last")
        )

        total_weight = float(project_order_df["Σύνολο kg"].sum())
        total_gunite = sum(item["result"].gunite_volume_m3 for item in st.session_state.project_elements)
        c1, c2, c3 = st.columns(3)
        with c1: card("Στοιχεία", str(len(elements_df)))
        with c2: card("Συνολικό gunite", f"{total_gunite:.3f} m³")
        with c3: card("Συνολικό βάρος χάλυβα/μανδυών", f"{total_weight:.2f} kg")

        st.markdown("### Στοιχεία έργου")
        st.dataframe(elements_df, width="stretch", hide_index=True)
        st.markdown("### Συγκεντρωτική παραγγελία")
        st.dataframe(summary_df, width="stretch", hide_index=True)
        with st.expander("Αναλυτική παραγγελία", expanded=False):
            st.dataframe(project_order_df, width="stretch", hide_index=True)

        project_input_rows = [input_dataframe(item["input"]).assign(Στοιχείο=item["input"].name) for item in st.session_state.project_elements]
        project_calc_rows = [calculation_dataframe(item["input"], item["result"]).assign(Στοιχείο=item["input"].name) for item in st.session_state.project_elements]
        project_inputs_df = pd.concat(project_input_rows, ignore_index=True) if project_input_rows else pd.DataFrame()
        project_calcs_df = pd.concat(project_calc_rows, ignore_index=True) if project_calc_rows else pd.DataFrame()
        excel_file = dataframe_to_excel_bytes({
            "Στοιχεία Έργου": elements_df, "Είσοδοι": project_inputs_df, "Υπολογισμοί": project_calcs_df,
            "Αναλυτική Παραγγελία": project_order_df, "Συγκεντρωτική": summary_df,
        }, title=f"GUNITE — {st.session_state.project_name}")
        c1, c2, c3 = st.columns(3)
        with c1:
            st.download_button("📥 Λήψη Excel", excel_file, file_name="paraggelia_ergou.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        with c2:
            technical_drawings = []
            for item in st.session_state.project_elements:
                try:
                    tech_fig = drawing.draw_technical_sheet(item["input"], item["result"], size="Μεγάλο")
                    technical_drawings.append(figure_to_png_bytes(tech_fig, dpi=220))
                    import matplotlib.pyplot as plt
                    plt.close(tech_fig)
                except Exception:
                    pass
            pdf_file = create_project_pdf(
                elements_df, project_order_df, summary_df,
                technical_drawings=technical_drawings,
            )
            st.download_button("📄 Λήψη PDF", pdf_file, file_name="anafora_gunite.pdf",
                               mime="application/pdf", use_container_width=True)
        with c3:
            if st.button("💾 Αποθήκευση έργου", type="primary", use_container_width=True):
                save_current_project()
                st.success("Το έργο αποθηκεύτηκε. Μπορείς να το ξανανοίξεις από την πλαϊνή στήλη.")
            if st.button("➕ Νέο στοιχείο", type="primary", use_container_width=True):
                reset_element()
                st.rerun()
